from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Create document
doc = Document()

# Set Landscape Orientation for "Slide" feel
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

# Helper to set RTL
def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

# Helper to add Slide Title
def add_slide_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.name = 'Arial'  # Or a standard Arabic font if available
    run.font.color.rgb = RGBColor(0, 51, 102) # Dark Blue
    set_rtl(p)
    # Add a decorative line
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run('________________________________________')
    run2.font.color.rgb = RGBColor(200, 200, 200)
    run2.font.size = Pt(14)
    # Add space
    doc.add_paragraph()

# Helper to add Bullet Point
def add_bullet(text):
    p = doc.add_paragraph()
    p.style = 'List Bullet'
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT # Arabic alignment
    run = p.add_run(text)
    run.font.size = Pt(20)
    run.font.name = 'Arial'
    set_rtl(p)

def add_normal_text(text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.size = Pt(18)
    run.font.name = 'Arial'
    if bold:
        run.font.bold = True
    set_rtl(p)

# ================= SLIDE 1: Title =================
# Vertically center title slide content (approximate with newlines)
for _ in range(3): doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_t = p_title.add_run('منظومة انجاز')
run_t.font.size = Pt(54)
run_t.font.bold = True
run_t.font.color.rgb = RGBColor(0, 51, 102)
set_rtl(p_title)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_s = p_sub.add_run('نظام إدارة وإصدار الشهادات الإشعاعية')
run_s.font.size = Pt(28)
run_s.font.color.rgb = RGBColor(100, 100, 100)
set_rtl(p_sub)

p_by = doc.add_paragraph()
p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_b = p_by.add_run('إعداد: م. إدريس فتح الله الهرى')
run_b.font.size = Pt(18)
set_rtl(p_by)

doc.add_page_break()

# ================= SLIDE 2: Def =================
add_slide_title('1. التعريف بمنظومة «انجاز»')
add_bullet('منظومة تقنية متكاملة لأتمتة إصدار الشهادات والتحاليل المخبرية.')
add_bullet('الغرض: الانتقال الكامل من الورقي إلى الرقمي.')
add_bullet('المشكلة: بطء الإجراءات، الأخطاء البشرية، صعوبة الأرشفة.')
add_bullet('الفئة المستهدفة: مركز القياسات الإشعاعية والتدريب.')
add_bullet('النطاق: من استلام العينة وحتى طباعة الشهادة والتقارير.')

doc.add_page_break()

# ================= SLIDE 3: Architecture =================
add_slide_title('2. هيكل المنظومة (System Architecture)')
add_bullet('نمط التصميم (MVVM): فصل الواجهة عن الكود لسهولة الصيانة.')
add_bullet('الواجهة (UI): تقنية WPF لتجربة مستخدم تفاعلية.')
add_bullet('المنطق (Backend): خدمات C# للحسابات والأمن.')
add_bullet('قاعدة البيانات: SQLite لأداء عالٍ وموثوقية محلية.')
add_bullet('الهيكلية: تصميم Modular يسمح بإضافة وحدات مستقبلية.')

doc.add_page_break()

# ================= SLIDE 4: Tech Stack =================
add_slide_title('3. التقنيات المستخدمة')
add_normal_text('تم استخدام أحدث تقنيات مايكروسوفت:', True)
add_bullet('لغة البرمجة: C# (.NET 8.0)')
add_bullet('واجهة المستخدم: WPF + Material Design')
add_bullet('توليد التقارير: QuestPDF (طباعة دقيقة)')
add_bullet('الرسوم البيانية: LiveCharts')
add_bullet('التشفير والباركود: QRCoder & ZXing')

doc.add_page_break()

# ================= SLIDE 5: Pros =================
add_slide_title('4. إيجابيات المنظومة')
add_bullet('أداء عالي وسرعة استجابة فائقة.')
add_bullet('واجهة عربية عصرية تدعم الوضع الليلي.')
add_bullet('أمان عالٍ للبيانات ونظام صلاحيات محكم.')
add_bullet('استقرار وموثوقية بفضل قاعدة بيانات SQLite.')
add_bullet('شهادات مؤمنة بباركود QR ضد التزوير.')

doc.add_page_break()

# ================= SLIDE 6: Cons =================
add_slide_title('5. التحديات الحالية (نقاط للتحسين)')
add_bullet('العمل محلياً فقط (تطبيق سطح مكتب).')
add_bullet('ارتباط ببيئة نظام Windows.')
add_bullet('عدم وجود ربط سحابي (Cloud) في النسخة الحالية.')
add_bullet('الحاجة لربط مباشر مع أجهزة المختبر (مستقبلاً).')

doc.add_page_break()

# ================= SLIDE 7: Roadmap =================
add_slide_title('6. سبل التطوير المستقبلية')
add_bullet('الانتقال للسحابة (Cloud Database).')
add_bullet('تطوير لوحة تحكم للهواتف الذكية.')
add_bullet('استخدام الذكاء الاصطناعي لتحليل البيانات التاريخية.')
add_bullet('أتمتة القراءة من أجهزة القياس مباشرة.')

doc.add_page_break()

# ================= SLIDE 8: Status =================
add_slide_title('7. الحالة والاختبارات')
add_bullet('تم إجراء اختبارات شاملة (الوظائف، الأداء، الأمان).')
add_bullet('النظام يعالج السيناريوهات المعقدة بكفاءة.')
add_bullet('جاهزية تامة للتشغيل الفعلي.')
add_bullet('تم التحقق من دقة الحسابات الإشعاعية.')

doc.add_page_break()

# ================= SLIDE 9: Conclusion =================
for _ in range(3): doc.add_paragraph()

p_end = doc.add_paragraph()
p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_e = p_end.add_run('شكراً لحسن استماعكم')
run_e.font.size = Pt(48)
run_e.font.bold = True
run_e.font.color.rgb = RGBColor(0, 102, 51)
set_rtl(p_end)

p_open = doc.add_paragraph()
p_open.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_o = p_open.add_run('المجال مفتوح للنقاش والاستفسارات')
run_o.font.size = Pt(24)
set_rtl(p_open)

# Save
output_path = r'd:\منظومة شهادات جديدة 2026 جديدة\عرض_منظومة_انجاز.docx'
doc.save(output_path)
print(f'Presentation saved to: {output_path}')
