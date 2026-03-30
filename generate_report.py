from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Create document
doc = Document()

# Set RTL for Arabic
def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def add_rtl_paragraph(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    set_rtl(p)
    return p

def add_rtl_heading(text, level):
    h = doc.add_heading(text, level)
    set_rtl(h)
    return h

# ===================== TITLE PAGE =====================
title = doc.add_heading('تقرير رسمي تفصيلي', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading('منظومة انجاز – Enjaz', 1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle2 = doc.add_heading('نظام إدارة وإصدار الشهادات والتحاليل الإشعاعية', 2)
subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
doc.add_paragraph('─' * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# Info section
info_para = doc.add_paragraph()
info_para.add_run('إعداد: ').bold = True
info_para.add_run('نظام الذكاء الاصطناعي التوليدي (Antigravity/Gemini)')
set_rtl(info_para)

info_para2 = doc.add_paragraph()
info_para2.add_run('بناءً على تعليمات: ').bold = True
info_para2.add_run('م. إدريس فتح الله الهرى')
set_rtl(info_para2)

info_para3 = doc.add_paragraph()
info_para3.add_run('تاريخ التقرير: ').bold = True
info_para3.add_run('30 ديسمبر 2025')
set_rtl(info_para3)

doc.add_paragraph()
doc.add_paragraph('─' * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

# ===================== SECTION 1 =====================
doc.add_page_break()
add_rtl_heading('1. مقدمة عامة', 1)

add_rtl_heading('1.1 تعريف موجز بالمنظومة', 2)
p = add_rtl_paragraph('منظومة «انجاز – Enjaz» هي منظومة رقمية متكاملة تم تطويرها لأتمتة عمليات إدارة وإصدار وتوثيق وأرشفة الشهادات الخاصة بالتحاليل والقياسات الإشعاعية. تهدف المنظومة إلى استبدال الإجراءات الورقية التقليدية بنظام رقمي مركزي يضمن الدقة وسهولة التتبع.')

add_rtl_heading('1.2 المجال الذي تخدمه المنظومة', 2)
add_rtl_paragraph('تخدم المنظومة قطاع الرقابة الإشعاعية، وتحديداً:')
bullets = [
    'إصدار شهادات خلو العينات من الإشعاع للمواد الاستهلاكية',
    'إصدار شهادات التحاليل الإشعاعية للعينات البيئية',
    'توثيق نتائج القياسات للنظائر المشعة (K40, Cs-137, Ra-226, Th-232, Raeq)',
    'أرشفة السجلات الإشعاعية للجهات الرقابية'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

add_rtl_heading('1.3 الهدف من إعداد هذا التقرير', 2)
add_rtl_paragraph('يهدف هذا التقرير إلى:')
bullets = [
    'تقديم وصف تقني وإداري شامل للمنظومة',
    'تمكين الإدارة العليا من تقييم المنظومة فنياً وإدارياً',
    'توثيق القدرات الحالية والتحديات والتوصيات المستقبلية'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

add_rtl_heading('1.4 الدور المتوقع في دعم التحول الرقمي', 2)
add_rtl_paragraph('يُتوقع أن تُسهم المنظومة في:')
bullets = [
    'تقليل الاعتماد على السجلات الورقية',
    'تسريع إجراءات إصدار الشهادات',
    'توحيد قالب الشهادات وضمان اتساقها',
    'تسهيل عمليات التدقيق والمراجعة',
    'دعم اتخاذ القرار عبر التقارير الإحصائية'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

# ===================== SECTION 2 =====================
doc.add_page_break()
add_rtl_heading('2. نبذة تعريفية عن منظومة «انجاز»', 1)

table = doc.add_table(rows=7, cols=2)
table.style = 'Table Grid'
data = [
    ('العنصر', 'التفاصيل'),
    ('الاسم الرسمي', 'انجاز – Enjaz'),
    ('الاسم التقني', 'CertificateSystem'),
    ('طبيعة المنظومة', 'منظومة رقمية متخصصة لإدارة الشهادات الإشعاعية'),
    ('نوع التطبيق', 'تطبيق سطح مكتب (Desktop Application)'),
    ('نطاق الاستخدام', 'داخلي / مؤسسي'),
    ('نظام التشغيل المستهدف', 'Microsoft Windows')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('2.1 الجهات المستفيدة', 2)
bullets = [
    'قسم التحاليل الإشعاعية: إدخال نتائج القياسات وإصدار الشهادات',
    'قسم الرقابة: متابعة الشهادات والتقارير الإحصائية',
    'الإدارة: الاطلاع على لوحة المعلومات والتقارير',
    'قسم الأرشيف: حفظ واسترجاع السجلات [غير مؤكد - حسب الهيكل التنظيمي الفعلي]'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

# ===================== SECTION 3 =====================
doc.add_page_break()
add_rtl_heading('3. تاريخ إنشاء المنظومة', 1)

table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('المرحلة', 'التاريخ'),
    ('بدء مرحلة التحليل والتصميم', '6 ديسمبر 2025'),
    ('بدء التطوير البرمجي', '7 ديسمبر 2025'),
    ('إطلاق النسخة التشغيلية الأولى', '20 ديسمبر 2025'),
    ('آخر تحديث موثق', '30 ديسمبر 2025')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('3.1 مراحل التطوير المسجلة', 2)
add_rtl_paragraph('بناءً على سجلات التطوير المتاحة، تضمنت التحديثات الأخيرة:')
bullets = [
    'تحسين نظام البحث وإضافة البحث برقم العينة',
    'تحسين نظام التقارير وإضافة أعمدة العينات البيئية والاستهلاكية',
    'إصلاح أخطاء تقنية في آلية الحفظ والتنقل',
    'تحسين تسجيل التعديلات (Audit Trail)',
    'إضافة دعم الترتيب التصاعدي للتقارير حسب التاريخ'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

# ===================== SECTION 4 =====================
doc.add_page_break()
add_rtl_heading('4. التقنيات المستخدمة في إنشاء المنظومة', 1)

add_rtl_heading('4.1 لغات البرمجة', 2)
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
data = [
    ('اللغة', 'الاستخدام'),
    ('C#', 'لغة البرمجة الأساسية'),
    ('XAML', 'تصميم واجهات المستخدم')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('4.2 الأطر البرمجية وبيئات التطوير', 2)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('الإطار/البيئة', 'الوصف'),
    ('.NET 8.0', 'إطار العمل الأساسي'),
    ('WPF (Windows Presentation Foundation)', 'إطار واجهة المستخدم'),
    ('Material Design in XAML Toolkit', 'مكتبة التصميم المرئي'),
    ('Microsoft.Extensions.DependencyInjection', 'إدارة التبعيات')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('4.3 قاعدة البيانات', 2)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('العنصر', 'التفاصيل'),
    ('نوع قاعدة البيانات', 'SQLite'),
    ('مكتبة الوصول', 'Microsoft.Data.Sqlite'),
    ('موقع التخزين', 'محلي على جهاز المستخدم'),
    ('السعة النظرية', 'ملايين السجلات (حتى 281 تيرابايت)')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('4.4 مكتبات إضافية', 2)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('المكتبة', 'الوظيفة'),
    ('QuestPDF', 'توليد ملفات PDF للشهادات والتقارير'),
    ('ClosedXML', 'تصدير التقارير إلى Excel'),
    ('LiveCharts.Wpf', 'الرسوم البيانية في لوحة المعلومات'),
    ('BCrypt.Net-Next', 'تشفير كلمات المرور')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('4.5 البنية المعتمدة', 2)
add_rtl_paragraph('النوع: تطبيق سطح مكتب مستقل (Standalone Desktop)')
add_rtl_paragraph('النمط المعماري: MVVM (Model-View-ViewModel)')
add_rtl_paragraph('طبقات التطبيق:')
bullets = [
    'طبقة العرض (Views - XAML)',
    'طبقة منطق العرض (ViewModels)',
    'طبقة الخدمات (Services)',
    'طبقة الوصول للبيانات (Repositories)',
    'طبقة النماذج (Models)'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

doc.add_paragraph()

add_rtl_heading('4.6 آليات الأمان والصلاحيات', 2)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('الآلية', 'الوصف'),
    ('تشفير كلمات المرور', 'خوارزمية BCrypt'),
    ('نظام الصلاحيات', 'صلاحيات متعددة المستويات (مدير، مستخدم عادي)'),
    ('سجل التدقيق', 'تسجيل جميع العمليات (إضافة، تعديل، حذف)'),
    ('التحقق من المستخدم', 'نظام تسجيل دخول إلزامي')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

# ===================== SECTION 5 =====================
doc.add_page_break()
add_rtl_heading('5. الجمهور المستهدف للمنظومة', 1)

add_rtl_heading('5.1 المستخدمون الفنيون', 2)
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
data = [
    ('الفئة', 'طبيعة الاستخدام'),
    ('أخصائيو القياسات الإشعاعية', 'إدخال نتائج التحاليل، إصدار الشهادات'),
    ('فنيو المختبرات', 'إدخال بيانات العينات')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('5.2 المستخدمون الإداريون', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('الفئة', 'طبيعة الاستخدام'),
    ('رؤساء الأقسام', 'مراجعة الشهادات، الموافقة، التقارير'),
    ('الإدارة العليا', 'لوحة المعلومات، التقارير الإحصائية'),
    ('مديرو النظام', 'إدارة المستخدمين، الإعدادات، النسخ الاحتياطي')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('5.3 الجهات الرقابية', 2)
add_rtl_paragraph('[غير مؤكد - يعتمد على طبيعة التكامل مع الجهات الخارجية]')

# ===================== SECTION 6 =====================
doc.add_page_break()
add_rtl_heading('6. مميزات منظومة «انجاز»', 1)

add_rtl_heading('6.1 مميزات تقنية', 2)
table = doc.add_table(rows=8, cols=2)
table.style = 'Table Grid'
data = [
    ('الميزة', 'الوصف'),
    ('أتمتة إصدار الشهادات', 'توليد رقم الشهادة تلقائياً بصيغة موحدة'),
    ('توليد PDF احترافي', 'شهادات جاهزة للطباعة بتنسيق ثابت'),
    ('الطباعة المباشرة', 'إمكانية الطباعة من داخل التطبيق'),
    ('البحث المتقدم', 'بحث بمعايير متعددة (رقم الشهادة، العينة، المورد، التاريخ)'),
    ('التصفح بالصفحات', 'عرض البيانات بشكل منظم مع دعم التصفح'),
    ('تقليل التدخل اليدوي', 'حسابات تلقائية، قيم افتراضية ذكية'),
    ('دعم الإكمال التلقائي', 'اقتراحات للحقول المتكررة')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('6.2 مميزات تنظيمية', 2)
table = doc.add_table(rows=7, cols=2)
table.style = 'Table Grid'
data = [
    ('الميزة', 'الوصف'),
    ('مركزية البيانات', 'قاعدة بيانات موحدة لجميع الشهادات'),
    ('سجل التدقيق الشامل', 'تتبع جميع التعديلات مع تفاصيل دقيقة'),
    ('التقارير الإحصائية', 'تقارير قابلة للتخصيص والتصدير'),
    ('لوحة معلومات تفاعلية', 'إحصائيات فورية ورسوم بيانية'),
    ('إدارة المستخدمين', 'تحكم كامل بالصلاحيات'),
    ('النسخ الاحتياطي', 'نظام نسخ احتياطي يدوي وتلقائي')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('6.3 مميزات واجهة المستخدم', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('الميزة', 'الوصف'),
    ('واجهة عربية كاملة', 'دعم كامل للغة العربية (RTL)'),
    ('تصميم عصري', 'واجهة Material Design'),
    ('سهولة الاستخدام', 'تصميم بديهي يقلل منحنى التعلم')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

# ===================== SECTION 7 =====================
doc.add_page_break()
add_rtl_heading('7. سلبيات المنظومة والتحديات الحالية', 1)

add_rtl_heading('7.1 القيود التقنية', 2)
table = doc.add_table(rows=4, cols=3)
table.style = 'Table Grid'
data = [
    ('القيد', 'الوصف', 'الأثر'),
    ('تطبيق محلي فقط', 'لا يدعم الوصول عبر الشبكة أو الإنترنت', 'يتطلب وجود المستخدم على نفس الجهاز'),
    ('قاعدة بيانات محلية', 'SQLite لا يدعم الوصول المتعدد المتزامن', 'قد يحد من العمل التشاركي'),
    ('نظام Windows فقط', 'لا يعمل على أنظمة أخرى', 'يتطلب أجهزة Windows')
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i].cells[j].text = cell_data
        if i == 0:
            for para in table.rows[i].cells[j].paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('7.2 التحديات التشغيلية', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('التحدي', 'الوصف'),
    ('النسخ الاحتياطي اليدوي', 'يعتمد على التزام المستخدم بالنسخ الدوري'),
    ('التدريب', 'يتطلب تدريب المستخدمين الجدد'),
    ('الصيانة', 'تحديثات البرنامج تتطلب تثبيتاً يدوياً')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('7.3 أوجه القصور المحتملة', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('الجانب', 'الملاحظة'),
    ('التكامل الخارجي', 'لا يوجد حالياً تكامل مع أنظمة خارجية'),
    ('الوصول عن بُعد', 'غير متاح حالياً'),
    ('التقارير المجدولة', 'تتطلب تشغيلاً يدوياً')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

# ===================== SECTION 8 =====================
doc.add_page_break()
add_rtl_heading('8. آليات مقترحة للتغلب على السلبيات', 1)

add_rtl_heading('8.1 حلول تقنية مقترحة [غير مؤكد - مقترحات فقط]', 2)
table = doc.add_table(rows=4, cols=3)
table.style = 'Table Grid'
data = [
    ('المقترح', 'الوصف', 'الأولوية'),
    ('تحويل إلى نظام ويب', 'إتاحة الوصول عبر المتصفح', 'طويلة المدى'),
    ('قاعدة بيانات مركزية', 'الانتقال إلى SQL Server أو PostgreSQL', 'متوسطة المدى'),
    ('واجهة برمجية (API)', 'لتمكين التكامل مع أنظمة أخرى', 'طويلة المدى')
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i].cells[j].text = cell_data
        if i == 0:
            for para in table.rows[i].cells[j].paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('8.2 حلول تنظيمية مقترحة [غير مؤكد - مقترحات فقط]', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('المقترح', 'الوصف'),
    ('جدول نسخ احتياطي', 'تحديد مسؤولية ومواعيد النسخ الاحتياطي'),
    ('دليل مستخدم', 'إعداد وثائق تدريبية شاملة'),
    ('خطة طوارئ', 'إجراءات استعادة البيانات')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('8.3 التمييز بين المطبق والمقترح', 2)
table = doc.add_table(rows=6, cols=2)
table.style = 'Table Grid'
data = [
    ('العنصر', 'الحالة'),
    ('النسخ الاحتياطي التلقائي', 'مطبق ✓'),
    ('سجل التدقيق', 'مطبق ✓'),
    ('التقارير القابلة للتخصيص', 'مطبق ✓'),
    ('الوصول عبر الشبكة', '[غير مؤكد - مقترح فقط]'),
    ('التكامل مع أنظمة خارجية', '[غير مؤكد - مقترح فقط]')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

# ===================== SECTION 9 =====================
doc.add_page_break()
add_rtl_heading('9. عناصر إضافية في توصيف المنظومة', 1)

add_rtl_heading('9.1 قابلية التوسع', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('الجانب', 'التقييم'),
    ('إضافة أنواع شهادات جديدة', 'ممكن عبر تطوير برمجي'),
    ('زيادة عدد المستخدمين', 'محدود بطبيعة التطبيق المحلي'),
    ('زيادة حجم البيانات', 'قابل للتوسع (ملايين السجلات)')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('9.2 قابلية التكامل', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('النظام', 'إمكانية التكامل'),
    ('أنظمة ERP', '[غير مؤكد - يتطلب تطوير API]'),
    ('أنظمة المراسلات', '[غير مؤكد - يتطلب تطوير]'),
    ('أنظمة الأرشفة', 'متاح عبر التصدير (PDF/Excel)')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('9.3 دعم الأرشفة والتدقيق', 2)
table = doc.add_table(rows=4, cols=2)
table.style = 'Table Grid'
data = [
    ('الميزة', 'التفاصيل'),
    ('سجل التدقيق', 'يسجل: المستخدم، التاريخ، نوع العملية، التفاصيل'),
    ('تتبع التعديلات', 'يحفظ تفاصيل كل تعديل على مستوى الحقول'),
    ('الأرشفة', 'تخزين دائم مع إمكانية البحث والاسترجاع')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('9.4 أثر المنظومة على جودة اتخاذ القرار', 2)
bullets = [
    'لوحة المعلومات: توفر رؤية فورية لحجم العمل والإحصائيات',
    'التقارير: تدعم تحليل الأداء والاتجاهات',
    'البيانات الموثقة: تقلل احتمالية الأخطاء في القرارات'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

# ===================== SECTION 10 =====================
doc.add_page_break()
add_rtl_heading('10. خاتمة وتوصيات للإدارة العليا', 1)

add_rtl_heading('10.1 خلاصة التقييم العام', 2)
add_rtl_paragraph('تُمثل منظومة «انجاز» حلاً تقنياً متكاملاً لإدارة وإصدار الشهادات الإشعاعية. تتميز المنظومة بواجهة عربية احترافية، وآليات أتمتة فعالة، ونظام تدقيق شامل. تعتمد على تقنيات حديثة ومستقرة، وتوفر مرونة في التقارير والتصدير.')

doc.add_paragraph()

add_rtl_heading('10.2 مدى جاهزية المنظومة للتشغيل المستدام', 2)
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
data = [
    ('الجانب', 'التقييم'),
    ('الاستقرار التقني', 'جاهز للتشغيل ✓'),
    ('اكتمال الوظائف الأساسية', 'مكتمل ✓'),
    ('الوثائق والتدريب', 'يحتاج تعزيز'),
    ('خطة الطوارئ', 'يحتاج توثيق رسمي')
]
for i, (c1, c2) in enumerate(data):
    table.rows[i].cells[0].text = c1
    table.rows[i].cells[1].text = c2
    if i == 0:
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

doc.add_paragraph()

add_rtl_heading('10.3 التوصيات', 2)

p = doc.add_paragraph()
p.add_run('توصيات فورية (0-3 أشهر):').bold = True
set_rtl(p)
bullets = [
    'اعتماد المنظومة رسمياً للتشغيل',
    'تدريب المستخدمين على كافة الوظائف',
    'تحديد مسؤولية النسخ الاحتياطي ومتابعتها',
    'إعداد دليل مستخدم مختصر'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

doc.add_paragraph()

p = doc.add_paragraph()
p.add_run('توصيات متوسطة المدى (3-12 شهر):').bold = True
set_rtl(p)
bullets = [
    'تقييم أداء المنظومة بعد فترة تشغيل كافية',
    'جمع ملاحظات المستخدمين للتحسين',
    'دراسة إمكانية التوسع للعمل الشبكي'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

doc.add_paragraph()

p = doc.add_paragraph()
p.add_run('توصيات طويلة المدى (12+ شهر) [غير مؤكد - مقترحات فقط]:').bold = True
set_rtl(p)
bullets = [
    'دراسة جدوى التحول إلى نظام ويب',
    'دراسة التكامل مع الأنظمة المؤسسية الأخرى',
    'تطوير تطبيق محمول للمتابعة'
]
for b in bullets:
    p = doc.add_paragraph(b, style='List Bullet')
    set_rtl(p)

# ===================== DISCLAIMER =====================
doc.add_page_break()
add_rtl_heading('تنويه هام', 1)
add_rtl_paragraph('هذا التقرير مُعد بناءً على المعلومات التقنية المتاحة من خلال التعامل المباشر مع الكود المصدري للمنظومة. أي عنصر موسوم بـ [غير مؤكد] يتطلب التحقق من الجهات المعنية قبل اعتماده.')

doc.add_paragraph()
doc.add_paragraph('─' * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

p = doc.add_paragraph('نهاية التقرير')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p.runs:
    run.bold = True
    run.font.size = Pt(14)

doc.add_paragraph()
doc.add_paragraph('─' * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

p = doc.add_paragraph('تم إعداد هذا التقرير بتاريخ 30 ديسمبر 2025')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p.runs:
    run.italic = True

p = doc.add_paragraph('بناءً على تعليمات م. إدريس فتح الله الهرى')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in p.runs:
    run.italic = True

# Save
doc.save(r'd:\منظومة شهادات جديدة 2026 جديدة\تقرير_منظومة_انجاز_كامل.docx')
print('تم حفظ التقرير الكامل بنجاح!')
print(r'الموقع: d:\منظومة شهادات جديدة 2026 جديدة\تقرير_منظومة_انجاز_كامل.docx')
